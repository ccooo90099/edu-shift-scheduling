"""把排班结果写成 Excel。

**明细表的列名与源数据保持一致**，所以输出可以直接喂回体检做闭环对照 ——
这是「工具排的比手排好多少」能被量化的前提。
"""
from __future__ import annotations

from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ...domain.model.period import Period

HEAD = PatternFill("solid", fgColor="E6EFF1")
BAD = PatternFill("solid", fgColor="FAE6E6")
WARN = PatternFill("solid", fgColor="FFF6E8")
THIN = Side(style="thin", color="C6D0D5")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)

DETAIL_COLUMNS = ["团队ID", "期位", "中心", "校区", "区域", "程度", "产品",
                  "团队类型", "团队名称", "首次服务时间", "主指导员",
                  "指导室", "是否促销"]


def _autosize(ws, limit=40):
    widths: dict[int, int] = defaultdict(int)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                longest = max(len(line) for line in str(cell.value).split("\n"))
                widths[cell.column] = max(widths[cell.column], longest)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(width + 3, limit)


def _header(ws, labels):
    ws.append(labels)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = BOX
    ws.freeze_panes = "A2"


def write_detail(wb, teams, season):
    ws = wb.create_sheet("排班明细")
    _header(ws, DETAIL_COLUMNS)
    for t in sorted(teams, key=lambda x: (x.center, x.grade.value, x.product.name)):
        ws.append([
            t.id,
            t.period.label(season) if t.period else "",
            t.center, t.campus, t.region, t.grade.value, t.product.name,
            t.tier.value, t.name,
            t.slot.text if t.slot else "",
            t.instructor or "", t.room or "",
            "是" if t.is_promotional else "否"])
        if not t.is_scheduled:
            for cell in ws[ws.max_row]:
                cell.fill = BAD
    _autosize(ws)


def write_board(wb, teams, season):
    """看板：行 = 校区，列 = 期位 × 时段。"""
    ws = wb.create_sheet("看板")
    placed = [t for t in teams if t.is_scheduled]
    if not placed:
        ws.append(["没有已排定的团队"])
        return
    periods = sorted({t.period for t in placed}, key=lambda p: p.value)
    slots = sorted({t.slot for t in placed}, key=lambda s: s.start)
    columns = [(p, s) for p in periods for s in slots]
    _header(ws, ["校区"] + ["%s\n%s" % (p.label(season), s.text) for p, s in columns])

    grid: dict[tuple, list[str]] = defaultdict(list)
    for t in placed:
        grid[(t.campus, t.period, t.slot)].append(
            "%s/%s/%s/%s" % (t.grade.value, t.tier.value,
                             t.product.name[:2], t.instructor or "?"))
    for campus in sorted({t.campus for t in placed}):
        ws.append([campus] + ["\n".join(sorted(grid.get((campus, p, s), [])))
                              for p, s in columns])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = BOX
    _autosize(ws, limit=28)


def write_instructor_day(wb, teams, season):
    ws = wb.create_sheet("老师课表")
    _header(ws, ["主指导员", "期位", "时段", "中心", "产品", "程度", "教室"])
    rows = [t for t in teams if t.is_scheduled]
    for t in sorted(rows, key=lambda x: (x.instructor, x.period.value, x.slot.start)):
        ws.append([t.instructor, t.period.label(season), t.slot.text,
                   t.center, t.product.name, t.grade.value, t.room or ""])
    _autosize(ws)


def write_issues(wb, report):
    """问题清单。**已验证**与**估算候选**分列，绝不混为一谈；
    因缺数据而未能检查的项单独一段列出。"""
    ws = wb.create_sheet("问题清单")
    _header(ws, ["规则", "结论强度", "说明"])
    for v in report.violations:
        ws.append([v.rule, v.label, v.detail])
        if not v.is_verified:
            for cell in ws[ws.max_row]:
                cell.fill = WARN
    if report.unchecked:
        ws.append([])
        ws.append(["—— 因缺数据而未能检查的项 ——"])
        ws[ws.max_row][0].font = Font(bold=True)
        for u in report.unchecked:
            ws.append(["", "未检查", u])
            for cell in ws[ws.max_row]:
                cell.fill = WARN
    _autosize(ws, limit=90)


def write_summary(wb, report, season, outcome=None):
    ws = wb.create_sheet("总览", 0)
    ws.append(["场景", season.value])
    if outcome is not None:
        ws.append(["求解状态", outcome.status])
        ws.append(["目标值", outcome.objective])
        ws.append(["下界", outcome.best_bound])
        ws.append(["用时（秒）", round(outcome.elapsed_seconds, 1)])
        for note in outcome.notes:
            ws.append(["提示", note])
    ws.append([])
    ws.append(["总分", report.score.total])
    for k, n in sorted(report.score.breakdown.items(),
                       key=lambda kv: -report.score.weights[kv[0]] * kv[1]):
        ws.append([k, n, "× %d" % report.score.weights[k],
                   n * report.score.weights[k]])
    ws.append([])
    ws.append(["已验证违规", report.verified_count])
    ws.append(["估算候选（需真实授课日核查）", report.candidate_count])
    ws.append(["未检查项", len(report.unchecked)])
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    _autosize(ws, limit=60)


def write_workbook(path, teams, report, season, outcome=None):
    wb = Workbook()
    wb.remove(wb.active)
    write_summary(wb, report, season, outcome)
    write_detail(wb, teams, season)
    write_board(wb, teams, season)
    write_instructor_day(wb, teams, season)
    write_issues(wb, report)
    wb.save(str(path))
    return path
