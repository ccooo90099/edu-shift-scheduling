"""把排班结果写成 Excel：明细表 + 看板 + 老师课表 + 问题清单。"""
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .slots import parse_slot

HEAD = PatternFill("solid", fgColor="E6EFF1")
BAND1 = PatternFill("solid", fgColor="D6EEF5")      # 段次1 那种青色
BAND2 = PatternFill("solid", fgColor="F7DCE6")      # 段次2 那种粉色
BAD = PatternFill("solid", fgColor="FAE6E6")
THIN = Side(style="thin", color="C6D0D5")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)


def _autosize(ws, limit=40):
    widths = defaultdict(int)
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                longest = max(len(line) for line in str(cell.value).split("\n"))
                widths[cell.column] = max(widths[cell.column], longest)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(width + 3, limit)


def _header(ws, labels):
    ws.append(labels)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD
        cell.border = BOX


def write_detail(wb, teams, result):
    ws = wb.create_sheet("排班明细")
    # 列名和源表保持一致，这样这张表能直接喂回 health_check ——
    # 排班和体检用同一把尺子，才谈得上"比原来好多少"
    _header(ws, ["团队ID", "段次", "周", "区域", "中心", "校区", "程度", "产品",
                 "团队类型", "团队名称", "首次服务时间", "主指导员", "指导室",
                 "是否促销", "原时段", "原主指导员", "是否改动"])
    for t in sorted(teams, key=lambda x: (x.段次, x.周, x.校区, x.程度, x.产品)):
        slot, inst = result.assignments.get(t.团队ID, ("", ""))
        changed = "" if not t.原时段 else (
            "否" if (slot, inst) == (t.原时段, t.原主指导员) else "是")
        ws.append([t.团队ID, t.段次, t.周, t.区域, t.中心, t.校区, t.程度, t.产品,
                   t.团队类型, t.团队名称, slot, inst, "", t.是否促销,
                   t.原时段, t.原主指导员, changed])
        if not slot:
            for cell in ws[ws.max_row]:
                cell.fill = BAD
    _autosize(ws)


def write_board(wb, teams, result, slots):
    """截图那张看板：行 = 校区，列 = 段次 × 时段。"""
    ws = wb.create_sheet("看板")
    slots = sorted(slots, key=lambda s: parse_slot(s)[0])
    段次列表 = sorted({t.段次 for t in teams})

    columns = [("校区", None, None)]
    for 段次 in 段次列表:
        for s in slots:
            columns.append((段次, s, BAND1 if 段次列表.index(段次) % 2 == 0 else BAND2))

    ws.append([c[0] for c in columns])
    ws.append([""] + [c[1] for c in columns[1:]])
    for row in (ws[1], ws[2]):
        for i, cell in enumerate(row):
            cell.font = Font(bold=True)
            cell.fill = columns[i][2] or HEAD
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center")

    cells = defaultdict(list)
    for t in teams:
        slot, inst = result.assignments.get(t.团队ID, ("", ""))
        if not slot:
            continue
        label = "%s/%s/%s/%s" % (t.程度, t.团队类型, t.产品, inst)
        cells[(t.校区, t.段次, slot)].append(label)

    for 校区 in sorted({t.校区 for t in teams}):
        row = [校区]
        for 段次, s, _ in columns[1:]:
            row.append("\n".join(sorted(cells.get((校区, 段次, s), []))))
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.border = BOX
            cell.alignment = WRAP
    _autosize(ws, limit=34)


def write_instructors(wb, teams, result, slots):
    ws = wb.create_sheet("老师课表")
    slots = sorted(slots, key=lambda s: parse_slot(s)[0])
    _header(ws, ["主指导员", "段次", "周", "课时数", "去几个中心"] + slots)

    by_key = defaultdict(dict)
    meta = defaultdict(set)
    team_of = {t.团队ID: t for t in teams}
    for tid, (slot, inst) in result.assignments.items():
        t = team_of[tid]
        by_key[(inst, t.段次, t.周)][slot] = "%s %s/%s" % (t.中心, t.产品, t.团队类型)
        meta[(inst, t.段次, t.周)].add(t.中心)

    for key in sorted(by_key):
        row = list(key) + [len(by_key[key]), len(meta[key])]
        row += [by_key[key].get(s, "") for s in slots]
        ws.append(row)
        if len(meta[key]) > 1:
            ws.cell(ws.max_row, 5).fill = BAD
    _autosize(ws, limit=30)


def write_problems(wb, teams, result):
    ws = wb.create_sheet("问题清单")
    _header(ws, ["类型", "对象", "说明"])
    team_of = {t.团队ID: t for t in teams}
    for tid in result.未排上:
        t = team_of.get(tid)
        ws.append(["没排上", tid,
                   "%s %s %s %s" % (t.中心, t.程度, t.产品, t.团队类型) if t else ""])
    for group, rule, strength in result.连堂未满足:
        ws.append(["连堂未满足（%s）" % strength, "·".join(map(str, group)), rule])
    if ws.max_row == 1:
        ws.append(["—", "—", "没有问题"])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BOX
    _autosize(ws)


def write_summary(wb, result):
    ws = wb.create_sheet("总览", 0)
    _header(ws, ["项", "值"])
    rows = [("状态", result.状态), ("罚分", round(result.罚分)),
            ("排上", len(result.assignments)), ("没排上", len(result.未排上)),
            ("连堂未满足", len(result.连堂未满足)),
            ("其中硬约束", sum(1 for x in result.连堂未满足 if x[2] == "硬")),
            ("用时秒", result.用时秒)]
    rows += [("罚分·" + k, round(v)) for k, v in sorted(result.分项.items()) if v]
    for item in rows:
        ws.append(list(item))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BOX
    _autosize(ws)


def write_workbook(path, teams, result, slots):
    wb = Workbook()
    wb.remove(wb.active)
    write_summary(wb, result)
    write_detail(wb, teams, result)
    write_board(wb, teams, result, slots)
    write_instructors(wb, teams, result, slots)
    write_problems(wb, teams, result)
    wb.save(path)
    return path
