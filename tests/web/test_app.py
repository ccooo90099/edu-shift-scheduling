"""网页层 —— 端到端。用 TestClient，不起真服务器。"""
import pytest
from fastapi.testclient import TestClient

from scheduling.application.dto import SolveTask, TaskStatus
from scheduling.domain.model.curriculum import product_of
from scheduling.domain.model.instructor import Instructor
from scheduling.domain.model.venue import Center, Datum
from scheduling.infrastructure.persistence.sqlite import (
    SqliteCalendarRepository, SqliteCenterRepository,
    SqliteInstructorRepository, SqliteTaskRepository, connect)
from scheduling.interfaces.web.app import create_app
from scheduling.interfaces.web.container import Container
from scheduling.interfaces.web.gate import COOKIE_NAME, PasswordGate


TEST_PASSWORD = "demo-pass"


@pytest.fixture
def app(tmp_path):
    conn = connect(tmp_path / "t.db")
    return create_app(
        Container(
            centers=SqliteCenterRepository(conn),
            instructors=SqliteInstructorRepository(conn),
            calendars=SqliteCalendarRepository(conn),
            tasks=SqliteTaskRepository(conn),
            data_dir=tmp_path),
        gate=PasswordGate(password=TEST_PASSWORD, secret="test-secret"))


@pytest.fixture
def anon(app):
    """没过门禁的客户端。"""
    return TestClient(app)


@pytest.fixture
def client(app):
    """过了门禁的客户端 —— 绝大多数测试用这个。"""
    c = TestClient(app)
    r = c.post("/login", data={"password": TEST_PASSWORD, "next": "/"})
    assert r.status_code in (200, 303)
    return c


def test_首页能打开(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "排班任务" in r.text


def test_没有数据时顶部就挂着告警条(client):
    """缺坐标、缺资质都不会报错，只会让求解悄悄退化 —— 所以要一直喊。"""
    r = client.get("/")
    assert "还没填学年日历" in r.text
    assert "不管的话" in r.text, "告警必须说清后果，否则没人去补"


def test_告警条说明后果而不只是报状态(client):
    r = client.get("/api/readiness")
    codes = {i["code"] for i in r.json()}
    assert "calendar.empty" in codes
    assert all(i["consequence"] for i in r.json())


def test_补齐数据后告警消失(client):
    client.post("/centers", data={"name": "甲中心", "rooms": "2"})
    assert "1 个中心还没有坐标" in client.get("/").text
    client.post("/centers/甲中心/coordinate",
                data={"longitude": "114.06", "latitude": "22.54",
                      "map_source": "高德"})
    assert "还没有坐标" not in client.get("/").text


def test_坐标按地图来源定坐标系(client, app):
    client.post("/centers", data={"name": "甲中心"})
    client.post("/centers/甲中心/coordinate",
                data={"longitude": "114.06", "latitude": "22.54",
                      "map_source": "百度"})
    c = app.state.container.centers.get("甲中心")
    assert c.coordinate.datum is Datum.BD09, "百度是 BD-09，不是 GCJ-02"


def test_未知地图来源被拒绝而不是猜一个(client):
    client.post("/centers", data={"name": "甲中心"})
    r = client.post("/centers/甲中心/coordinate",
                    data={"longitude": "114.06", "latitude": "22.54",
                          "map_source": "某某地图"})
    assert r.status_code == 400
    assert "未知地图来源" in r.json()["detail"]


def test_中心页可按缺坐标筛选(client):
    client.post("/centers", data={"name": "有坐标"})
    client.post("/centers", data={"name": "没坐标"})
    client.post("/centers/有坐标/coordinate",
                data={"longitude": "114.0", "latitude": "22.5",
                      "map_source": "天地图"})
    body = client.get("/centers?filter=missing-coordinate").text
    assert "没坐标" in body and body.count("有坐标") <= 1


def test_指导员单个与批量编辑(client, app):
    client.post("/instructors", data={
        "name": "甲老师", "products": "编程理论,双语文化",
        "unavailable": "18:30-20:30", "max_sessions": "3"})
    i = app.state.container.instructors.get("甲老师")
    assert product_of("双语文化") in i.teachable
    assert i.max_sessions_per_day == 3

    client.post("/instructors/bulk",
                data={"names": "乙老师,丙老师", "products": "文学美育",
                      "max_sessions": "5"})
    for name in ("乙老师", "丙老师"):
        x = app.state.container.instructors.get(name)
        assert x.teachable == {product_of("文学美育")}
        assert x.max_sessions_per_day == 5


def test_未知产品被拒绝(client):
    r = client.post("/instructors", data={"name": "甲", "products": "不存在的课"})
    assert r.status_code == 400 and "未知产品" in r.json()["detail"]


def test_缺配置的指导员能筛出来(client):
    client.post("/instructors", data={"name": "配了的", "products": "编程理论"})
    client.post("/instructors", data={"name": "没配的"})
    body = client.get("/instructors?filter=unconfigured").text
    assert "没配的" in body


def test_建任务后跳到任务页(client):
    r = client.post("/tasks", data={"name": "段次3", "season": "寒暑假"},
                    follow_redirects=True)
    assert r.status_code == 200 and "段次3" in r.text


def test_进度片段在未结束时带轮询_结束后不带(client, app):
    """任务终结后返回的片段不再带 hx-* ，前端自然停下来。"""
    tasks = app.state.container.tasks
    t = SolveTask(id="t1", name="x", season="寒暑假", status=TaskStatus.RUNNING)
    tasks.create(t)
    assert "hx-trigger" in client.get("/tasks/t1/progress").text

    t.status = TaskStatus.DONE
    tasks.update(t)
    assert "hx-trigger" not in client.get("/tasks/t1/progress").text


def test_任务状态与求解状态分开显示(client, app):
    """任务「已完成」不代表解「最优」—— 把 FEASIBLE 说成排好了是误导。"""
    app.state.container.tasks.create(SolveTask(
        id="t2", name="x", season="寒暑假", status=TaskStatus.DONE,
        solver_status="FEASIBLE", objective=100.0, best_bound=80.0))
    data = client.get("/api/tasks/t2").json()
    assert data["status_label"] == "已完成"
    assert data["solver_status"] == "FEASIBLE"
    assert "未证明最优" in data["gap_hint"]

    body = client.get("/tasks/t2").text
    assert "FEASIBLE" in body and "未证明最优" in body


def test_不存在的任务返回404而不是500(client):
    assert client.get("/tasks/nope").status_code == 404
    assert client.get("/api/tasks/nope").status_code == 404


def test_没有产物时下载返回404(client, app):
    app.state.container.tasks.create(
        SolveTask(id="t3", name="x", season="寒暑假"))
    assert client.get("/tasks/t3/download").status_code == 404


# ── 教室数可配（用户补充）───────────────────────────────────

def test_教室数可以随时改_不只是新建时(client, app):
    """校区的教室是有限且已知的，后台要能直接填。"""
    client.post("/centers", data={"name": "甲中心", "rooms": "3"})
    assert app.state.container.centers.get("甲中心").room_count == 3

    client.post("/centers/甲中心/rooms", data={"count": "6"})
    c = app.state.container.centers.get("甲中心")
    assert c.room_count == 6
    assert not c.rooms_are_estimated, "人填的是实数，不是估计"


def test_留空教室数时不动原值(client, app):
    client.post("/centers", data={"name": "甲中心", "rooms": "5"})
    client.post("/centers", data={"name": "甲中心", "region": "福田"})
    c = app.state.container.centers.get("甲中心")
    assert c.room_count == 5, "只改区域不该把教室数清零"
    assert c.region == "福田"


def test_教室数填负数被拒(client):
    client.post("/centers", data={"name": "甲中心"})
    assert client.post("/centers/甲中心/rooms",
                       data={"count": "-1"}).status_code == 400


# ── 区域相邻（用户补充）─────────────────────────────────────

def test_区域相邻页列出所有区域并标出未配的(client):
    client.post("/centers", data={"name": "宝安甲", "region": "宝安"})
    client.post("/centers", data={"name": "南山甲", "region": "南山"})
    body = client.get("/regions").text
    assert "宝安" in body and "南山" in body
    assert "未配，一律按跨区域算" in body


def test_配了相邻关系后自动补成双向(client):
    client.post("/centers", data={"name": "宝安甲", "region": "宝安"})
    client.post("/centers", data={"name": "南山甲", "region": "南山"})
    client.post("/regions", data={"region": "宝安", "neighbours": "南山"})

    body = client.get("/regions").text
    assert "未配" not in body, "两边都该显示已配 —— 只写一边系统自动补反向"


def test_相邻表落盘后重启仍在(client, app, tmp_path):
    client.post("/centers", data={"name": "宝安甲", "region": "宝安"})
    client.post("/regions", data={"region": "宝安", "neighbours": "南山,福田"})
    assert (tmp_path / "region_adjacency.json").exists()

    from scheduling.interfaces.web.app import _load_adjacency
    assert _load_adjacency(app.state.container)["宝安"] == ["南山", "福田"]


# ── 场地容量两层（用户补充：教室数不一定，做成配置）──────────

def test_校区容量页列出各校区及其中心(client):
    client.post("/centers", data={"name": "百花科学", "campus": "百花", "rooms": "3"})
    client.post("/centers", data={"name": "百花文学", "campus": "百花", "rooms": "2"})
    body = client.get("/capacity").text
    assert "百花" in body
    assert "百花科学（3 间）" in body and "百花文学（2 间）" in body
    assert "只按各中心教室数" in body, "没配校区上限时要说明"


def test_设了校区上限后标出哪一层更严(client):
    client.post("/centers", data={"name": "百花科学", "campus": "百花", "rooms": "3"})
    client.post("/centers", data={"name": "百花文学", "campus": "百花", "rooms": "2"})
    client.post("/capacity", data={"campus": "百花", "cap": "2"})
    body = client.get("/capacity").text
    assert "校区上限更严" in body


def test_校区上限留空表示不设这一层(client, app):
    from scheduling.interfaces.web.app import _load_caps
    client.post("/centers", data={"name": "甲中心", "campus": "甲", "rooms": "3"})
    client.post("/capacity", data={"campus": "甲", "cap": "2"})
    assert _load_caps(app.state.container) == {"甲": 2}

    client.post("/capacity", data={"campus": "甲", "cap": ""})
    assert _load_caps(app.state.container) == {}, "留空 = 清掉这一层，不是设成 0"


def test_校区上限非法输入被拒(client):
    client.post("/centers", data={"name": "甲中心", "campus": "甲"})
    assert client.post("/capacity", data={"campus": "甲", "cap": "abc"}).status_code == 400
    assert client.post("/capacity", data={"campus": "甲", "cap": "-1"}).status_code == 400


# ── 口令门禁 ────────────────────────────────────────────────
#
# ⚠️ 这是门禁不是认证：全站一个共享口令、不分人、泄露了只能换一个重启。
# 它挡的是误闯的路人，不是保护真实数据的手段。

def test_没口令时页面被挡去登录页(anon):
    r = anon.get("/", follow_redirects=False)
    assert r.status_code == 303
    assert r.headers["location"].startswith("/login")


def test_没口令时API返回401而不是一坨登录页HTML(anon):
    """curl 拿到 HTML 会很莫名，而且状态码还是 200。"""
    r = anon.get("/api/readiness")
    assert r.status_code == 401
    assert r.json()["detail"].endswith("/login")


def test_登录页和静态资源不需要口令(anon):
    assert anon.get("/login").status_code == 200
    assert anon.get("/static/app.css").status_code == 200


def test_探活不需要口令(anon):
    """否则容器会一直被判成不健康 —— 不能拿 /api/readiness 当探活。"""
    r = anon.get("/healthz")
    assert r.status_code == 200 and r.text == "ok"


def test_口令对了才放行(anon):
    assert anon.post("/login", data={"password": "错的"},
                     follow_redirects=False).headers["location"].endswith("error=1")
    assert anon.get("/", follow_redirects=False).status_code == 303

    anon.post("/login", data={"password": TEST_PASSWORD})
    assert anon.get("/").status_code == 200


def test_登录后回到原来想去的页面(anon):
    r = anon.post("/login", data={"password": TEST_PASSWORD, "next": "/centers"},
                  follow_redirects=False)
    assert r.headers["location"] == "/centers"


def test_不做开放重定向(anon):
    """next=//evil.com 会被浏览器当成协议相对 URL 跳到外站。"""
    for bad in ("//evil.com", "https://evil.com", "javascript:alert(1)"):
        r = anon.post("/login", data={"password": TEST_PASSWORD, "next": bad},
                      follow_redirects=False)
        assert r.headers["location"] == "/", bad


def test_通行证不能伪造(anon):
    anon.cookies.set(COOKIE_NAME, "99999999999.deadbeef")
    assert anon.get("/", follow_redirects=False).status_code == 303


def test_通行证会过期(app):
    """签发时就把过期时间签进去，改了签名就对不上。"""
    g = app.state.gate
    token = g.issue(now=0)
    assert g.verify(token, now=0)
    assert not g.verify(token, now=g.ttl_seconds + 1)


def test_cookie是httponly(anon):
    r = anon.post("/login", data={"password": TEST_PASSWORD},
                  follow_redirects=False)
    assert "httponly" in r.headers["set-cookie"].lower()


def test_退出后又被挡在外面(client):
    assert client.get("/").status_code == 200
    client.post("/logout")
    assert client.get("/", follow_redirects=False).status_code == 303


def test_登录页把门禁的局限说清楚(anon):
    """别让人以为这是安全模块。"""
    body = anon.get("/login").text
    assert "不是认证" in body
    assert "共享口令" in body


def test_随机口令不含易混字符():
    """0/O、1/l/I 口头念或手抄容易错。"""
    from scheduling.interfaces.web.gate import generate_password
    for _ in range(50):
        assert not (set(generate_password()) & set("0O1lI"))


# ── 种子数据（免费 PaaS 重启清空文件系统）───────────────────

def test_库空时灌入示例数据_已有数据时不动(tmp_path):
    from scheduling.application.use_cases.seed_demo import SeedDemo
    from scheduling.infrastructure.persistence.sqlite import connect

    conn = connect(tmp_path / "s.db")
    centers = SqliteCenterRepository(conn)
    instructors = SqliteInstructorRepository(conn)
    calendars = SqliteCalendarRepository(conn)
    seed = SeedDemo(centers, instructors, calendars)

    assert seed() is True
    n = len(centers.all())
    assert n >= 8 and len(instructors.all()) >= 20
    assert len(calendars.load()) == 2

    assert seed() is False, "已有数据就不该再灌"
    assert len(centers.all()) == n


def test_示例数据不含真实门店或姓名():
    """真实门店信息与 106 位指导员姓名一路都没进过仓库，这里不破例。"""
    from scheduling.application.use_cases.demo_dataset import DemoDataset
    d = DemoDataset()
    centers = d.centers()
    assert all(c.name.startswith("示范") for c in centers)
    assert all(i.name.startswith("示教") for i in d.instructors(centers))
    for c in centers:
        assert 113.7 < c.coordinate.longitude < 114.7
        assert 22.4 < c.coordinate.latitude < 22.9, "坐标在深圳范围内即可"


def test_示例数据照搬了真实表的两条硬结构():
    """不是随便编的：文学楼只上文学美育（真实表 26/26 行），
    S7 不开物理化学。结构不像，跑出来的结果就没有参考价值。"""
    from scheduling.application.use_cases.demo_dataset import DemoDataset
    d = DemoDataset()
    df = d.teams_frame(d.centers())
    文学楼 = df[df["中心"].str.endswith("文学")]
    assert set(文学楼["产品"]) == {"文学美育"}
    assert not ({"躬行实践", "溯源"} & set(df[df["程度"] == "S7"]["产品"]))
    assert {"躬行实践"} <= set(df[df["程度"] == "S8"]["产品"])


def test_示例数据能直接拿去求解(tmp_path):
    """灌完就该是个能跑的系统，而不是一堆填了一半的配置。"""
    from scheduling.application.use_cases.seed_demo import SeedDemo
    from scheduling.infrastructure.persistence.sqlite import connect

    conn = connect(tmp_path / "s.db")
    centers = SqliteCenterRepository(conn)
    instructors = SqliteInstructorRepository(conn)
    SeedDemo(centers, instructors, SqliteCalendarRepository(conn))()

    assert all(c.has_coordinate for c in centers.all()), "都有坐标，转场才不是粗判"
    assert all(c.room_count > 0 for c in centers.all())
    assert all(i.is_configured for i in instructors.all()), "都配了资质"
    assert not centers.without_coordinate()
    assert not instructors.unconfigured()


# ── 学年日历（线上报的 /calendar 404）────────────────────────

def test_导航和告警条里的链接都不是死链(client):
    """/calendar 曾经在导航和告警条里都指着，但路由根本没实现，
    线上点进去 404。这条守住：凡是页面里出现的站内链接都必须打得开。"""
    import re
    seen = set()
    for page in ("/", "/centers", "/capacity", "/regions", "/instructors",
                 "/calendar"):
        body = client.get(page).text
        seen |= set(re.findall(r'href="(/[^"#?]*)', body))
    broken = [u for u in sorted(seen)
              if not u.startswith("/static") and client.get(u).status_code != 200]
    assert not broken, "死链：%s" % broken


def test_日历页能增删批次(client, app):
    assert client.get("/calendar").status_code == 200
    client.post("/calendar", data={
        "name": "测试批次", "season": "春秋季",
        "start": "2026-09-05", "end": "2026-12-20",
        "dates_a": "2026-09-05,2026-09-12", "dates_b": "2026-09-06"})
    cal = app.state.container.calendars.load()
    assert [b.name for b in cal] == ["测试批次"]
    from scheduling.domain.model.period import Period
    assert len(list(cal)[0].dates_of(Period.A)) == 2

    client.post("/calendar/测试批次/delete")
    assert len(app.state.container.calendars.load()) == 0


def test_同名批次视为覆盖而不是重复添加(client, app):
    for end in ("2026-12-20", "2026-12-31"):
        client.post("/calendar", data={"name": "同名", "season": "春秋季",
                                       "start": "2026-09-05", "end": end})
    cal = list(app.state.container.calendars.load())
    assert len(cal) == 1 and cal[0].end.isoformat() == "2026-12-31"


def test_日期写错要报400而不是静默吞掉(client):
    r = client.post("/calendar", data={"name": "x", "season": "春秋季",
                                       "start": "2026-09-05", "end": "2026-12-20",
                                       "dates_a": "九月五号"})
    assert r.status_code == 400 and "2026-07-06" in r.json()["detail"]


# ── htmx 缺失时的兜底 ───────────────────────────────────────

def test_没有htmx时进度靠整页刷新兜底(client, app):
    """线上 htmx 没加载上，进度就永远停在「求解中」。
    现在这件事不再取决于那个静态文件在不在。"""
    from scheduling.application.dto import SolveTask, TaskStatus
    app.state.container.tasks.create(
        SolveTask(id="tf", name="x", season="寒暑假", status=TaskStatus.RUNNING))

    page = client.get("/tasks/tf").text
    assert "!window.htmx" in page, "htmx 不在时才插 meta refresh，在的话不插"
    assert "httpEquiv" in page

    t = app.state.container.tasks.get("tf")
    t.status = TaskStatus.DONE
    app.state.container.tasks.update(t)
    assert "httpEquiv" not in client.get("/tasks/tf").text, "跑完了就别再刷"


def test_htmx缺失提示说的是降级而不是坏了(client):
    body = client.get("/").text
    assert "已降级为整页刷新" in body
    assert "功能都能用" in body


# ── 上传 → 求解 → 看结果（这条链之前是断的）─────────────────

def _清单(tmp_path, centers, n_products=3):
    import pandas as pd
    产品 = ["编程理论", "双语文化", "文学美育"][:n_products]
    rows = [{"团队ID": "T%d%d" % (ci, pi), "段次": "段次1", "周": "星期六",
             "区域": "福田区", "中心": center, "程度": "S8", "产品": p,
             "团队类型": "LI", "团队名称": "%s%s" % (center, p),
             "首次服务时间": "08:10-10:10", "主指导员": "", "指导室": "01",
             "是否促销": "否"}
            for ci, center in enumerate(centers) for pi, p in enumerate(产品)]
    path = tmp_path / "清单.xlsx"
    pd.DataFrame(rows).to_excel(path, index=False)
    return path


def _等任务(tasks, tid, seconds=90):
    import time
    for _ in range(seconds):
        t = tasks.get(tid)
        if t.status.is_terminal:
            return t
        time.sleep(1)
    return tasks.get(tid)


@pytest.fixture
def seeded(app):
    from scheduling.application.use_cases.seed_demo import SeedDemo
    c = app.state.container
    SeedDemo(c.centers, c.instructors, c.calendars)()
    return c


def test_上传后会真的开始求解而不是永远排队(client, seeded, tmp_path):
    """这条链之前是断的：create_task 只存了一条记录就返回，
    没人调用求解器，任务永远停在「排队中」。"""
    from scheduling.application.use_cases.demo_dataset import DemoDataset
    src = _清单(tmp_path, [c.name for c in DemoDataset().centers()[:3]])
    r = client.post("/tasks", data={"name": "e2e", "season": "寒暑假"},
                    files={"upload": ("清单.xlsx", src.read_bytes(),
                                      "application/vnd.ms-excel")},
                    follow_redirects=False)
    tid = r.headers["location"].rsplit("/", 1)[-1]

    task = _等任务(seeded.tasks, tid)
    assert task.status.value == "done", task.error
    assert task.solver_status in ("OPTIMAL", "FEASIBLE")
    assert task.output_path, "要写出 xlsx，下载才有东西"
    assert task.score_breakdown, "要有得分明细"


def test_没传文件的任务直接标失败而不是卡在排队中(client, seeded):
    r = client.post("/tasks", data={"name": "空的", "season": "寒暑假"},
                    follow_redirects=False)
    tid = r.headers["location"].rsplit("/", 1)[-1]
    task = seeded.tasks.get(tid)
    assert task.status.value == "failed"
    assert "没有上传团队清单" in task.error


def test_一键跑示例_不用准备任何文件就能看到结果(client, seeded):
    """打开就能看结果 —— 没有现成的表也不该卡住。"""
    import re
    r = client.post("/tasks/demo", follow_redirects=False)
    assert r.status_code == 303
    tid = r.headers["location"].rsplit("/", 1)[-1]

    task = _等任务(seeded.tasks, tid, seconds=180)
    assert task.status.value == "done", task.error
    assert task.output_path

    body = client.get("/tasks/%s/board" % tid).text
    cells = re.findall(r'<span class="cellitem">([^<]+)</span>', body)
    assert cells, "看板要有格子"
    assert all(c.count("/") == 3 for c in cells), "格式是 程度/层级/科目/老师"


def test_下载的是五个sheet的完整工作簿(client, seeded, tmp_path):
    from openpyxl import load_workbook
    r = client.post("/tasks/demo", follow_redirects=False)
    tid = r.headers["location"].rsplit("/", 1)[-1]
    assert _等任务(seeded.tasks, tid, seconds=180).status.value == "done"

    resp = client.get("/tasks/%s/download" % tid)
    assert resp.status_code == 200
    out = tmp_path / "dl.xlsx"
    out.write_bytes(resp.content)
    assert set(load_workbook(out).sheetnames) == {
        "总览", "排班明细", "看板", "老师课表", "问题清单"}


def test_配置里的旧权重键报错要给出改名指引(client):
    """光说「不认识」没用，得告诉人改成什么。"""
    from scheduling.domain.policy.weights import Weights
    with pytest.raises(ValueError) as e:
        Weights.from_config({"跨中心一次": 800, "瞎写的": 1})
    msg = str(e.value)
    assert "现在叫「转场一次」" in msg
    assert "已废弃" in msg
    assert "可用的键" in msg


def test_跑完后课表直接显示在任务页上_不用再点一次(client, seeded):
    """用户说「结果呢」—— 原先任务页只有状态和得分，课表要再点一层。
    结果本身才是主角。"""
    import re
    tid = client.post("/tasks/demo", follow_redirects=False) \
        .headers["location"].rsplit("/", 1)[-1]
    assert _等任务(seeded.tasks, tid, seconds=180).status.value == "done"

    html = client.get("/tasks/%s" % tid).text
    assert "排课结果" in html
    assert 'class="board"' in html, "看板要直接嵌在任务页上"
    assert re.findall(r'<span class="cellitem">', html), "要有排课格子"
    assert "下载 Excel" in html


def test_没排上的团队要列出来而不是只给个数字(client, seeded):
    """「6 个没排上」没用 —— 得知道是哪 6 个，才谈得上加老师还是加教室。"""
    tid = client.post("/tasks/demo", follow_redirects=False) \
        .headers["location"].rsplit("/", 1)[-1]
    task = _等任务(seeded.tasks, tid, seconds=180)
    if not task.unplaced:
        pytest.skip("这次全排上了，没有可列的")

    html = client.get("/tasks/%s" % tid).text
    assert "没排上的 %d 个团队" % task.unplaced in html
    assert "加能教这门课的老师" in html, "要说清怎么才能排进去"


def test_总分要解释清楚而不是甩一个数字(client, seeded):
    """「目标值 600240」对人没有意义，得说明它主要由什么构成。"""
    tid = client.post("/tasks/demo", follow_redirects=False) \
        .headers["location"].rsplit("/", 1)[-1]
    task = _等任务(seeded.tasks, tid, seconds=180)
    html = client.get("/tasks/%s" % tid).text
    if task.unplaced:
        assert "没排上的权重远高于其他" in html


def test_跑完但没有产物时页面不会是空的(client, app):
    """条件写成「没跑完才显示进度卡」的话，跑完却没产物时两张卡都不显示。"""
    from scheduling.application.dto import SolveTask, TaskStatus
    app.state.container.tasks.create(SolveTask(
        id="noout", name="x", season="寒暑假", status=TaskStatus.DONE,
        solver_status="FEASIBLE", objective=1.0, best_bound=0.0))
    html = client.get("/tasks/noout").text
    assert "求解进度" in html and "FEASIBLE" in html
