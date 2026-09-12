"""端到端：造一份 xlsx → 跑 CLI → 验产物。

这是 engine/ 迁移到 src/scheduling/ 之后，证明整条链没断的测试。
"""
import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

时段 = ["08:10-10:10", "10:30-12:30", "14:00-16:00", "16:20-18:20"]
产品 = ["编程理论", "双语文化", "文学美育"]


@pytest.fixture
def 明细(tmp_path):
    """两个中心 × 三个产品，教室号完整，带主指导员。"""
    rows = []
    for ci, center in enumerate(["百花科学", "石厦科学"]):
        for pi, product in enumerate(产品):
            rows.append({
                "团队ID": "T%d%d" % (ci, pi), "段次": "段次1", "周": "星期六",
                "区域": "福田区", "中心": center, "程度": "S8", "产品": product,
                "团队类型": "LI", "团队名称": "%s%s1" % (center, product),
                "首次服务时间": 时段[pi], "主指导员": "老师%d%d" % (ci, pi),
                "指导室": "0%d" % (pi + 1), "是否促销": "否"})
    path = tmp_path / "明细.xlsx"
    pd.DataFrame(rows).to_excel(path, index=False)
    return path


def 跑(module, *args):
    return subprocess.run(
        [sys.executable, str(ROOT / "tools" / module), *map(str, args)],
        cwd=ROOT, capture_output=True, encoding="utf-8", errors="replace",
        timeout=180)


def test_体检跑得通并如实报告未检查项(明细):
    r = 跑("health_check.py", 明细, "--config", "config/rules.example.yaml")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "总分" in r.stdout
    # 没有学年日历 → 跨批次撞车必须显式报「未检查」，不能不声不响跳过
    assert "未检查" in r.stdout
    assert "跨批次" in r.stdout


def test_排班跑得通并写出五个sheet(明细, tmp_path):
    out = tmp_path / "结果.xlsx"
    r = 跑("schedule.py", 明细, "--config", "config/rules.example.yaml",
           "--out", out, "--time-limit", "20")
    assert r.returncode == 0, r.stdout + r.stderr
    assert out.exists()

    wb = load_workbook(out)
    assert set(wb.sheetnames) == {"总览", "排班明细", "看板", "老师课表", "问题清单"}

    明细页 = wb["排班明细"]
    assert 明细页.max_row == 7, "6 个团队 + 表头"
    时段列 = [明细页.cell(r, 10).value for r in range(2, 8)]
    assert all(时段列), "每个团队都该排上时段"


def test_排班会打印实际生效的权重(明细, tmp_path):
    """曾连续三轮调权重完全不生效，因为配置里还是旧值。"""
    r = 跑("schedule.py", 明细, "--config", "config/rules.example.yaml",
           "--out", tmp_path / "x.xlsx", "--time-limit", "10")
    assert "实际生效的权重" in r.stdout
    assert "排不上团队" in r.stdout


def test_求解状态被如实打印而不是笼统说排好了(明细, tmp_path):
    r = 跑("schedule.py", 明细, "--config", "config/rules.example.yaml",
           "--out", tmp_path / "x.xlsx", "--time-limit", "10")
    assert "求解状态" in r.stdout
    assert any(s in r.stdout for s in ("OPTIMAL", "FEASIBLE"))
    if "FEASIBLE" in r.stdout and "OPTIMAL" not in r.stdout:
        assert "未证明最优" in r.stdout


def test_排班结果能回灌体检形成闭环(明细, tmp_path):
    """明细页的列名与源表一致，所以产物可以直接当输入再体检一遍 ——
    这是「工具比手排好多少」能被量化的前提。"""
    out = tmp_path / "结果.xlsx"
    assert 跑("schedule.py", 明细, "--config", "config/rules.example.yaml",
              "--out", out, "--time-limit", "20").returncode == 0

    回灌 = tmp_path / "回灌.xlsx"
    pd.read_excel(out, sheet_name="排班明细").to_excel(回灌, index=False)
    r = 跑("health_check.py", 回灌, "--config", "config/rules.example.yaml")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "总分" in r.stdout


def test_没有配置文件也能跑(明细, tmp_path):
    """配置缺失时退回默认时段表，不该直接崩。"""
    r = 跑("schedule.py", 明细, "--config", "不存在.yaml",
           "--out", tmp_path / "x.xlsx", "--time-limit", "10")
    assert r.returncode == 0, r.stdout + r.stderr
